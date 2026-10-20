"""Bounded read-only extraction; document text and code are never executed."""
from pathlib import Path
import csv,io,json,re,zipfile
from . import config
from .scanner import safe_file_path
TEXT_EXTS={".txt",".md",".py",".js",".ts",".html",".css",".json",".xml",".yaml",".yml",".toml",".java",".c",".h",".cpp",".go",".rs",".sh",".sql"}
def clip(text): return (text or "")[:config.MAX_TEXT_CHARS]
def decode(data):
 for enc in ("utf-8-sig","utf-8","cp1252","latin-1"):
  try: return data.decode(enc)
  except UnicodeDecodeError: pass
 return data.decode("utf-8",errors="replace")
def extract_file(relative_path):
 path=safe_file_path(relative_path)
 if path.stat().st_size>config.MAX_FILE_BYTES: raise ValueError("File exceeds configured size limit.")
 ext=path.suffix.lower()
 if ext in TEXT_EXTS: return {"kind":"text","text":clip(decode(path.read_bytes())),"details":{"format":ext,"execution":"never"}}
 if ext==".rtf":
  raw=decode(path.read_bytes())
  try:
   from striprtf.striprtf import rtf_to_text
   text=rtf_to_text(raw)
  except ImportError: text=re.sub(r"\\'[0-9a-fA-F]{2}|\\[a-zA-Z]+-?\d* ?|[{}]"," ",raw)
  return {"kind":"text","text":clip(text),"details":{"format":"RTF"}}
 if ext==".csv":
  raw=decode(path.read_bytes())
  try: dialect=csv.Sniffer().sniff(raw[:8192],delimiters=",;\t|")
  except csv.Error: dialect=csv.excel
  rows=csv.reader(io.StringIO(raw),dialect); preview=[]; count=0
  for row in rows:
   count+=1
   if len(preview)<16: preview.append([str(v)[:250] for v in row[:24]])
   if count>=20000: break
  headers=preview[0] if preview else []
  return {"kind":"spreadsheet","text":clip(json.dumps({"headers":headers,"preview_rows":preview[:8],"sampled_rows":count,"count_capped_at":20000},ensure_ascii=False)),"details":{"format":"CSV","sampled_rows":count,"columns":len(headers)}}
 if ext==".pdf":
  from pypdf import PdfReader
  reader=PdfReader(str(path),strict=False)
  if reader.is_encrypted: raise ValueError("Encrypted PDFs are not opened.")
  chunks=[]
  for i,page in enumerate(reader.pages[:50]):
   chunks.append(f"[Page {i+1}]\n{page.extract_text() or ''}")
   if sum(map(len,chunks))>=config.MAX_TEXT_CHARS: break
  return {"kind":"document","text":clip("\n\n".join(chunks)),"details":{"format":"PDF","pages_total":len(reader.pages),"pages_read":len(chunks)}}
 if ext==".docx":
  from docx import Document
  doc=Document(str(path)); chunks=[p.text for p in doc.paragraphs if p.text.strip()]
  for table in doc.tables[:20]: chunks.append("[Table]\n"+"\n".join(" | ".join(c.text[:200] for c in row.cells[:20]) for row in table.rows[:100]))
  return {"kind":"document","text":clip("\n".join(chunks)),"details":{"format":"DOCX","paragraphs":len(doc.paragraphs),"tables":len(doc.tables)}}
 if ext==".xlsx":
  from openpyxl import load_workbook
  wb=load_workbook(str(path),read_only=True,data_only=True); sheets=[]
  try:
   for ws in wb.worksheets[:20]:
    rows=[[str(v)[:150] if v is not None else "" for v in row] for row in ws.iter_rows(max_row=150,max_col=30,values_only=True)]
    sheets.append({"sheet":ws.title,"reported_rows":ws.max_row,"reported_columns":ws.max_column,"preview_rows":rows[:10]})
  finally: wb.close()
  return {"kind":"spreadsheet","text":clip(json.dumps({"sheets":sheets},ensure_ascii=False,default=str)),"details":{"format":"XLSX","sheets":len(sheets),"read_only":True}}
 if ext==".xls":
  import xlrd
  book=xlrd.open_workbook(str(path),on_demand=True); sheets=[]
  try:
   for sh in book.sheets()[:20]:
    rows=[[str(sh.cell_value(r,c))[:150] for c in range(min(sh.ncols,30))] for r in range(min(sh.nrows,100))]
    sheets.append({"sheet":sh.name,"rows":sh.nrows,"columns":sh.ncols,"preview_rows":rows[:10]})
  finally: book.release_resources()
  return {"kind":"spreadsheet","text":clip(json.dumps({"sheets":sheets},ensure_ascii=False)),"details":{"format":"XLS legacy","sheets":len(sheets),"read_only":True}}
 if ext==".pptx":
  from pptx import Presentation
  prs=Presentation(str(path)); chunks=[]
  for i,slide in enumerate(prs.slides[:80]):
   lines=[s.text for s in slide.shapes if getattr(s,"has_text_frame",False) and s.text.strip()]
   chunks.append(f"[Slide {i+1}]\n"+"\n".join(lines))
   if sum(map(len,chunks))>=config.MAX_TEXT_CHARS: break
  return {"kind":"presentation","text":clip("\n\n".join(chunks)),"details":{"format":"PPTX","slides_total":len(prs.slides),"slides_read":len(chunks)}}
 if ext in {".png",".jpg",".jpeg",".webp"}:
  from PIL import Image
  with Image.open(path) as im: info={"format":im.format,"width":im.width,"height":im.height,"mode":im.mode,"ocr":"not enabled; metadata only"}
  return {"kind":"image","text":json.dumps(info),"details":info}
 if ext==".zip":
  items=[]
  with zipfile.ZipFile(path) as archive:
   for info in archive.infolist()[:500]:
    if ((info.external_attr>>16)&0o170000)==0o120000: continue
    name=info.filename.replace("\\","/")
    if name.startswith("/") or ".." in name.split("/"): name="[unsafe path omitted]"
    items.append({"name":name[:240],"size":info.file_size,"directory":info.is_dir()})
  return {"kind":"archive_inventory","text":clip(json.dumps(items,ensure_ascii=False)),"details":{"format":"ZIP inventory only","entries_listed":len(items),"extracted":False}}
 raise ValueError("This file type is not supported for content inspection.")
